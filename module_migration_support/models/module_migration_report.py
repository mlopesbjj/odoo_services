import base64
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from odoo import _, fields, models
from odoo.exceptions import UserError


HEADERS = [
    "name",
    "display_name",
    "installed_version",
    "latest_version",
    "author",
    "category",
    "application",
    "auto_install",
    "state",
]


class ModuleMigrationReport(models.Model):
    _name = "module.migration.report"
    _description = "Module Migration Report"
    _order = "create_date desc, id desc"

    name = fields.Char(
        required=True,
        default=lambda self: _("Module Migration Report %s")
        % fields.Datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )
    source_file = fields.Binary(string="Source XLSX")
    source_filename = fields.Char(string="Source Filename")
    export_file = fields.Binary(string="Current DB Modules XLSX", readonly=True)
    export_filename = fields.Char(string="Export Filename", readonly=True)
    result_file = fields.Binary(string="Comparison XLSX", readonly=True)
    result_filename = fields.Char(string="Result Filename", readonly=True)
    source_module_count = fields.Integer(string="Source Modules", readonly=True)
    target_module_count = fields.Integer(string="Current DB Modules", readonly=True)
    missing_module_count = fields.Integer(string="Missing in Current DB", readonly=True)
    common_module_count = fields.Integer(string="Common Modules", readonly=True)
    line_ids = fields.One2many(
        "module.migration.report.line",
        "report_id",
        string="Missing Modules",
        readonly=True,
    )

    def action_export_current_modules(self):
        for record in self:
            modules = record._get_installed_modules()
            content = record._build_export_workbook(modules)
            filename = "installed_modules_%s.xlsx" % fields.Date.today()
            record.write(
                {
                    "export_file": base64.b64encode(content),
                    "export_filename": filename,
                    "target_module_count": len(modules),
                }
            )
        return True

    def action_compare_modules(self):
        for record in self:
            if not record.source_file:
                raise UserError(_("Upload the source XLSX before running the comparison."))

            source_modules = record._load_modules_from_xlsx()
            target_modules = record._get_installed_modules()
            comparison = record._prepare_comparison(source_modules, target_modules)
            report_content = record._build_comparison_workbook(
                source_modules,
                target_modules,
                comparison["missing_modules"],
                comparison["common_modules"],
            )

            line_commands = [(5, 0, 0)]
            for module in comparison["missing_modules"]:
                line_commands.append(
                    (
                        0,
                        0,
                        {
                            "module_name": module["name"],
                            "display_name": module["display_name"],
                            "installed_version": module["installed_version"],
                            "latest_version": module["latest_version"],
                            "author": module["author"],
                            "category": module["category"],
                        },
                    )
                )

            filename = "module_migration_report_%s.xlsx" % fields.Date.today()
            record.write(
                {
                    "result_file": base64.b64encode(report_content),
                    "result_filename": filename,
                    "source_module_count": len(source_modules),
                    "target_module_count": len(target_modules),
                    "missing_module_count": len(comparison["missing_modules"]),
                    "common_module_count": len(comparison["common_modules"]),
                    "line_ids": line_commands,
                }
            )
        return True

    def action_open_export_file(self):
        self.ensure_one()
        if not self.export_file:
            raise UserError(_("Generate the current database export first."))
        return self._download_action("export_file", self.export_filename)

    def action_open_result_file(self):
        self.ensure_one()
        if not self.result_file:
            raise UserError(_("Generate the comparison file first."))
        return self._download_action("result_file", self.result_filename)

    def _download_action(self, field_name, filename):
        self.ensure_one()
        return {
            "type": "ir.actions.act_url",
            "url": "/web/content/%s/%s/%s?download=true"
            % (self._name, self.id, field_name),
            "target": "self",
        }

    def _get_installed_modules(self):
        modules = self.env["ir.module.module"].search(
            [("state", "=", "installed")], order="name"
        )
        return [self._serialize_module(module) for module in modules]

    def _serialize_module(self, module):
        return {
            "name": module.name or "",
            "display_name": module.shortdesc or "",
            "installed_version": module.installed_version or "",
            "latest_version": module.latest_version or "",
            "author": module.author or "",
            "category": module.category_id.name or "",
            "application": bool(module.application),
            "auto_install": bool(module.auto_install),
            "state": module.state or "",
        }

    def _load_modules_from_xlsx(self):
        self.ensure_one()
        workbook = load_workbook(
            filename=BytesIO(base64.b64decode(self.source_file)),
            data_only=True,
        )
        if "source_modules" in workbook.sheetnames:
            sheet = workbook["source_modules"]
        elif "modules" in workbook.sheetnames:
            sheet = workbook["modules"]
        else:
            raise UserError(
                _(
                    "The source file must contain a 'modules' or 'source_modules' sheet."
                )
            )

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(value or "").strip() for value in rows[0]]
        modules = []
        for row in rows[1:]:
            if not any(row):
                continue
            values = {headers[index]: row[index] for index in range(len(headers))}
            module_data = {header: values.get(header, "") for header in HEADERS}
            modules.append(module_data)
        return sorted(modules, key=lambda item: str(item["name"]))

    def _prepare_comparison(self, source_modules, target_modules):
        source_by_name = {module["name"]: module for module in source_modules}
        target_by_name = {module["name"]: module for module in target_modules}

        missing_names = sorted(set(source_by_name) - set(target_by_name))
        common_names = sorted(set(source_by_name) & set(target_by_name))

        missing_modules = [source_by_name[name] for name in missing_names]
        common_modules = [
            {
                "name": name,
                "display_name": source_by_name[name].get("display_name", ""),
                "installed_version": source_by_name[name].get("installed_version", ""),
                "latest_version": target_by_name[name].get("latest_version", ""),
                "author": source_by_name[name].get("author", ""),
                "category": source_by_name[name].get("category", ""),
                "application": source_by_name[name].get("application", ""),
                "auto_install": source_by_name[name].get("auto_install", ""),
                "state": "installed_in_both",
            }
            for name in common_names
        ]
        return {
            "missing_modules": missing_modules,
            "common_modules": common_modules,
        }

    def _build_export_workbook(self, modules):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "modules"
        self._write_sheet(sheet, modules)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _build_comparison_workbook(
        self, source_modules, target_modules, missing_modules, common_modules
    ):
        workbook = Workbook()

        summary = workbook.active
        summary.title = "summary"
        summary["A1"] = "Metric"
        summary["B1"] = "Value"
        summary["A1"].font = Font(bold=True)
        summary["B1"].font = Font(bold=True)
        summary.append(["Source installed modules", len(source_modules)])
        summary.append(["Current DB installed modules", len(target_modules)])
        summary.append(["Missing in current DB", len(missing_modules)])
        summary.append(["Common modules", len(common_modules)])
        self._autosize(summary, ["A", "B"])

        source_sheet = workbook.create_sheet("source_modules")
        self._write_sheet(source_sheet, source_modules)

        target_sheet = workbook.create_sheet("target_modules")
        self._write_sheet(target_sheet, target_modules)

        missing_sheet = workbook.create_sheet("missing_in_target")
        self._write_sheet(missing_sheet, missing_modules)

        common_sheet = workbook.create_sheet("common_modules")
        self._write_sheet(common_sheet, common_modules)

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _write_sheet(self, sheet, modules):
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for module in modules:
            sheet.append([module.get(header, "") for header in HEADERS])
        self._autosize(sheet)

    def _autosize(self, sheet, columns=None):
        if columns is None:
            columns = [chr(ord("A") + index) for index in range(len(HEADERS))]
        for column in columns:
            max_length = 0
            for cell in sheet[column]:
                value = "" if cell.value is None else str(cell.value)
                if len(value) > max_length:
                    max_length = len(value)
            sheet.column_dimensions[column].width = min(max_length + 2, 60)


class ModuleMigrationReportLine(models.Model):
    _name = "module.migration.report.line"
    _description = "Module Migration Report Line"
    _order = "module_name"

    report_id = fields.Many2one(
        "module.migration.report",
        required=True,
        ondelete="cascade",
    )
    module_name = fields.Char(string="Technical Name", required=True, readonly=True)
    display_name = fields.Char(string="Display Name", readonly=True)
    installed_version = fields.Char(string="Source Installed Version", readonly=True)
    latest_version = fields.Char(string="Source Latest Version", readonly=True)
    author = fields.Char(string="Author", readonly=True)
    category = fields.Char(string="Category", readonly=True)
